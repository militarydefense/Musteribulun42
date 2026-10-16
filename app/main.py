"""Silah Sektörü Müşteri Bulucu – web sunucusu (FastAPI)."""
import csv
import io
import json
import asyncio
import hashlib
import hmac
import os
import secrets
import socket
import time
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from . import comtrade, db, jobs
from .gtip import EMBARGO_COUNTRIES, GTIP_CODES, keywords_for

ROOT = Path(__file__).resolve().parent.parent
STATIC = ROOT / "static"
COUNTRIES = json.loads((STATIC / "countries.json").read_text("utf-8"))
jobs.COUNTRY_NAMES.update({c["code"]: c["en"] for c in COUNTRIES})
TR_NAMES = {c["code"]: c["tr"] for c in COUNTRIES}

# ------------------------------------------------------------------ oturum / giriş
SESSION_DAYS = 30
PUBLIC_PATHS = {"/login", "/manifest.webmanifest", "/sw.js", "/static/style.css", "/static/icon-192.png",
                "/static/icon-512.png", "/static/apple-touch-icon.png", "/static/icon.svg"}


def panel_password() -> str:
    return os.environ.get("PANEL_PASSWORD") or db.get_setting("panel_password")


def _secret() -> bytes:
    s = db.get_setting("session_secret")
    if not s:
        s = secrets.token_hex(32)
        db.set_settings({"session_secret": s})
    return s.encode()


def make_token() -> str:
    exp = str(int(time.time()) + SESSION_DAYS * 86400)
    pw_tag = hashlib.sha256(panel_password().encode()).hexdigest()[:12]  # şifre değişince oturumlar düşer
    sig = hmac.new(_secret(), f"{exp}.{pw_tag}".encode(), hashlib.sha256).hexdigest()
    return f"{exp}.{sig}"


def valid_token(tok: str) -> bool:
    try:
        exp, sig = tok.split(".", 1)
        if int(exp) < time.time():
            return False
    except (ValueError, AttributeError):
        return False
    pw_tag = hashlib.sha256(panel_password().encode()).hexdigest()[:12]
    good = hmac.new(_secret(), f"{exp}.{pw_tag}".encode(), hashlib.sha256).hexdigest()
    return hmac.compare_digest(sig, good)


def ensure_password():
    """İnternete açık modda şifre yoksa rastgele bir şifre üretip konsola yazar."""
    if os.environ.get("APP_PUBLIC") == "1" and not panel_password():
        pw = secrets.token_urlsafe(9)
        db.set_settings({"panel_password": pw})
        print("\n" + "=" * 60 + f"\n  PANEL ŞİFRESİ OLUŞTURULDU: {pw}\n  (Ayarlar sekmesinden değiştirebilirsiniz)\n" + "=" * 60 + "\n",
              flush=True)


app = FastAPI(title="Silah Sektörü Müşteri Bulucu")


@app.middleware("http")
async def auth_middleware(request: Request, call_next):
    path = request.url.path
    if not panel_password() or path in PUBLIC_PATHS or valid_token(request.cookies.get("oturum", "")):
        return await call_next(request)
    if path.startswith("/api/"):
        return JSONResponse({"detail": "Oturum süresi doldu, lütfen tekrar giriş yapın."}, status_code=401)
    return RedirectResponse("/login")


@app.get("/login")
def login_page():
    return FileResponse(STATIC / "login.html")


@app.post("/login")
async def login(request: Request):
    form = await request.form()
    pw = str(form.get("password", ""))
    if not panel_password() or not secrets.compare_digest(pw.encode(), panel_password().encode()):
        await asyncio.sleep(1.5)  # kaba kuvvet denemelerini yavaşlat
        return RedirectResponse("/login?hata=1", status_code=303)
    resp = RedirectResponse("/", status_code=303)
    https = request.headers.get("x-forwarded-proto", request.url.scheme) == "https"
    resp.set_cookie("oturum", make_token(), max_age=SESSION_DAYS * 86400, httponly=True, samesite="lax",
                    secure=https)
    return resp


@app.get("/logout")
def logout():
    resp = RedirectResponse("/login", status_code=303)
    resp.delete_cookie("oturum")
    return resp


@app.get("/manifest.webmanifest")
def manifest():
    return FileResponse(STATIC / "manifest.webmanifest", media_type="application/manifest+json")


@app.get("/sw.js")
def service_worker():
    return FileResponse(STATIC / "sw.js", media_type="application/javascript",
                        headers={"Service-Worker-Allowed": "/", "Cache-Control": "no-cache"})


@app.get("/api/access")
def access_info(request: Request):
    """Telefondan bağlanmak için yerel ağ adreslerini döndürür."""
    port = request.url.port or (443 if request.url.scheme == "https" else 80)
    ips = set()
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ips.add(s.getsockname()[0])
        s.close()
    except OSError:
        pass
    try:
        for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET):
            ips.add(info[4][0])
    except OSError:
        pass
    lan = [f"http://{ip}:{port}" for ip in sorted(ips) if not ip.startswith("127.")]
    return {"current": str(request.base_url).rstrip("/"), "lan": lan, "password_set": bool(panel_password()),
            "public": os.environ.get("APP_PUBLIC") == "1"}


ensure_password()
app.mount("/static", StaticFiles(directory=STATIC), name="static")


@app.get("/")
def index():
    return FileResponse(STATIC / "index.html")


# ------------------------------------------------------------------ meta
@app.get("/api/meta")
def meta():
    return {"countries": COUNTRIES, "gtip": GTIP_CODES, "embargo": EMBARGO_COUNTRIES,
            "statuses": ["yeni", "iletisime_gecildi", "teklif_verildi", "musteri", "uygun_degil"]}


@app.get("/api/keywords")
def keywords(gtip: str = "9303", country: str = ""):
    return keywords_for(gtip, country)


# ------------------------------------------------------------------ ayarlar
HIDDEN_KEYS = {"session_secret"}
SECRET_KEYS = {"google_api_key", "youtube_api_key", "search_api_key", "meta_access_token",
               "comtrade_api_key", "panel_password"}


@app.get("/api/settings")
def get_settings():
    s = db.all_settings()
    out = {}
    for k, v in s.items():
        if k in HIDDEN_KEYS:
            continue
        out[k] = ("•" * 8 + v[-4:]) if k in SECRET_KEYS and v else v
    return out


@app.post("/api/settings")
async def save_settings(req: Request):
    data = await req.json()
    clean = {k: str(v).strip() for k, v in data.items()
             if not str(v).startswith("••••••••") and k not in HIDDEN_KEYS}
    if os.environ.get("APP_PUBLIC") == "1" and clean.get("panel_password") == "":
        clean.pop("panel_password")  # internete açık modda şifre kaldırılamaz
    db.set_settings(clean)
    return {"ok": True}


# ------------------------------------------------------------------ istatistik
@app.get("/api/stats")
def stats():
    s = db.stats()
    for r in s["by_country"]:
        r["country_tr"] = TR_NAMES.get(r["country_code"], r["country"])
    return s


@app.get("/api/map")
def map_points(country: str = ""):
    return db.map_points(country)


# ------------------------------------------------------------------ müşteriler
@app.get("/api/leads")
def leads(q: str = "", country: str = "", source: str = "", status: str = "", has_email: bool = False,
          has_website: bool = False, sort: str = "created_at", order: str = "desc", page: int = 1,
          per_page: int = 50):
    per_page = max(1, min(per_page, 500))
    total, rows = db.query_leads(q, country, source, status, has_email, has_website, sort, order,
                                 per_page, (page - 1) * per_page)
    for r in rows:
        r["country_tr"] = TR_NAMES.get(r.get("country_code") or "", r.get("country") or "")
        r["embargo"] = EMBARGO_COUNTRIES.get(r.get("country_code") or "")
    return {"total": total, "rows": rows}


@app.get("/api/leads/{lead_id}")
def lead(lead_id: int):
    r = db.get_lead(lead_id)
    if not r:
        raise HTTPException(404, "Kayıt bulunamadı")
    return r


@app.patch("/api/leads/{lead_id}")
async def patch_lead(lead_id: int, req: Request):
    db.update_lead(lead_id, await req.json())
    return db.get_lead(lead_id)


@app.post("/api/leads/delete")
async def delete(req: Request):
    db.delete_leads((await req.json()).get("ids", []))
    return {"ok": True}


EXPORT_COLS = [("name", "Firma"), ("country_tr", "Ülke"), ("city", "Şehir"), ("address", "Adres"),
               ("intl_phone", "Telefon"), ("emails", "E-posta"), ("website", "Web sitesi"),
               ("facebook", "Facebook"), ("instagram", "Instagram"), ("youtube", "YouTube"),
               ("linkedin", "LinkedIn"), ("whatsapp", "WhatsApp"), ("rating", "Puan"), ("reviews", "Yorum"),
               ("followers", "Takipçi"), ("source", "Kaynak"), ("maps_url", "Google Maps"),
               ("status", "Durum"), ("notes", "Notlar"), ("embargo", "Ambargo uyarısı")]


def _export_rows(**filters):
    _, rows = db.query_leads(**filters, limit=None)
    for r in rows:
        soc = r.get("socials") or {}
        r["country_tr"] = TR_NAMES.get(r.get("country_code") or "", r.get("country") or "")
        r["intl_phone"] = r.get("intl_phone") or r.get("phone") or ", ".join(r.get("phones_extra") or [])
        r["emails"] = ", ".join(r.get("emails") or [])
        r["embargo"] = EMBARGO_COUNTRIES.get(r.get("country_code") or "", "")
        for k in ("facebook", "instagram", "youtube", "linkedin", "whatsapp"):
            r[k] = soc.get(k, "")
        yield [r.get(k) if r.get(k) is not None else "" for k, _ in EXPORT_COLS]


@app.get("/api/export")
def export(fmt: str = "xlsx", q: str = "", country: str = "", source: str = "", status: str = "",
           has_email: bool = False, has_website: bool = False):
    f = dict(q=q, country=country, source=source, status=status, has_email=has_email, has_website=has_website)
    if fmt == "csv":
        buf = io.StringIO()
        buf.write("﻿")
        w = csv.writer(buf, delimiter=";")
        w.writerow([h for _, h in EXPORT_COLS])
        w.writerows(_export_rows(**f))
        return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv; charset=utf-8",
                                 headers={"Content-Disposition": "attachment; filename=musteriler.csv"})
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Müşteriler"
    ws.append([h for _, h in EXPORT_COLS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for row in _export_rows(**f):
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGHIJKLMNOPQRST", (34, 16, 16, 40, 18, 36, 30, 30, 30, 30, 30, 20,
                                                   8, 8, 10, 12, 30, 14, 30, 18)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=musteriler.xlsx"})


# ------------------------------------------------------------------ işler
@app.post("/api/jobs/{kind}")
async def start_job(kind: str, req: Request):
    if kind not in ("maps", "enrich", "youtube", "social"):
        raise HTTPException(400, "Geçersiz iş türü")
    p = await req.json()
    need = {"maps": "google_api_key", "social": "search_api_key"}.get(kind)
    if need and not db.get_setting(need):
        raise HTTPException(400, "Gerekli API anahtarı tanımlı değil. Lütfen Ayarlar sekmesini doldurun.")
    if kind == "youtube" and not (db.get_setting("youtube_api_key") or db.get_setting("google_api_key")):
        raise HTTPException(400, "YouTube Data API anahtarı tanımlı değil (Ayarlar).")
    names = ", ".join(TR_NAMES.get(c, c) for c in (p.get("countries") or [])[:4])
    if len(p.get("countries") or []) > 4:
        names += f" +{len(p['countries']) - 4}"
    titles = {"maps": "Google Maps taraması", "enrich": "E-posta / sosyal medya zenginleştirme",
              "youtube": "YouTube kanal araştırması", "social": "Facebook / Instagram araştırması"}
    title = f"{titles[kind]}" + (f" – {names}" if names else "")
    return {"id": jobs.submit(kind, title, p)}


@app.get("/api/jobs")
def list_jobs():
    return db.list_jobs()


@app.get("/api/jobs/{job_id}")
def job(job_id: int):
    j = db.get_job(job_id)
    if not j:
        raise HTTPException(404)
    return j


@app.post("/api/jobs/{job_id}/cancel")
def cancel(job_id: int):
    jobs.cancel(job_id)
    return {"ok": True}


# ------------------------------------------------------------------ GTİP pazar analizi
_cache = {}


@app.get("/api/gtip/market")
def gtip_market(code: str = "9303", year: int = 2023):
    key = (code, year)
    if key not in _cache:
        try:
            rows = comtrade.market_analysis(code, year, db.get_setting("comtrade_api_key"))
        except Exception as e:  # noqa: BLE001
            raise HTTPException(502, f"UN Comtrade verisine ulaşılamadı: {e}")
        for r in rows:
            r["embargo"] = EMBARGO_COUNTRIES.get(r["iso2"], "")
        _cache[key] = rows
    rows = _cache[key]
    counts = {r["country_code"]: r["n"] for r in db.country_counts()}
    return [{**r, "leads": counts.get(r["iso2"], 0)} for r in rows]
